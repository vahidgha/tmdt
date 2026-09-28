"""
خروجی Excel و PDF مشترک برای همه گزارش‌ها و لیست‌های پنل.

طراحی: یک تابع واحد برای هر فرمت که همه‌جا استفاده می‌شود — نه پیاده‌سازی
جداگانه در هر route. هر دو خروجی یک سربرگ یکسان دارند: نام سایت، عنوان
گزارش، تاریخ و ساعت تولید (شمسی) و نام کاربر تولیدکننده.
"""
import io
import os
from datetime import datetime

from flask import send_file
from flask_login import current_user

from .jalali import gregorian_to_jalali

NAVY = "0E1B32"
GOLD = "B8925A"


class ExportDependencyError(RuntimeError):
    """پکیج‌های لازم برای خروجی (reportlab/arabic-reshaper/python-bidi) نصب نیستند."""


def _site_title() -> str:
    from .. import db_session
    from ..models import SiteSetting
    row = db_session.get(SiteSetting, "site_title")
    return row.value if row else "سامانه مدیریت اعضا و امور مالی"


def _now_jalali_str() -> str:
    now = datetime.now()
    jy, jm, jd = gregorian_to_jalali(now.year, now.month, now.day)
    return f"{jy}/{jm:02d}/{jd:02d} — {now.strftime('%H:%M')}"


def _generated_by() -> str:
    try:
        if current_user and current_user.is_authenticated:
            return current_user.full_name
    except Exception:
        pass
    return "—"


# ─────────────────────────────────────────────────────────────────────────────
# Excel
# ─────────────────────────────────────────────────────────────────────────────

def export_xlsx(headers, rows, report_title: str, filename: str, subtitle: str = ""):
    """
    یک فایل اکسل تمیز و راست‌به‌چپ با سربرگ اطلاعاتی می‌سازد.

    headers: لیست عنوان ستون‌ها
    rows:    لیست ردیف‌ها (هر ردیف یک لیست هم‌طول با headers)
    report_title: عنوان گزارش (مثلاً «گزارش مالی و اقساط»)
    subtitle: توضیح فیلترهای اعمال‌شده (اختیاری) — مثلاً «پروژه: ثمین — وضعیت: معوق»
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    n_cols = max(1, len(headers))
    wb = Workbook()
    ws = wb.active
    ws.title = report_title[:31] or "گزارش"
    ws.sheet_view.rightToLeft = True

    navy_fill = PatternFill("solid", fgColor=NAVY)
    gold_font = Font(bold=True, color=GOLD, size=16)
    white_font = Font(bold=True, color="FFFFFF", size=10)
    muted_font = Font(color="6B7280", size=10)
    thin = Side(style="thin", color="E6E7EB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_cursor = 1

    def _merged(text, font, height=22, fill=None):
        nonlocal row_cursor
        ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=n_cols)
        c = ws.cell(row=row_cursor, column=1, value=text)
        c.font = font
        c.alignment = Alignment(horizontal="center", vertical="center")
        if fill:
            for ci in range(1, n_cols + 1):
                ws.cell(row=row_cursor, column=ci).fill = fill
        ws.row_dimensions[row_cursor].height = height
        row_cursor += 1

    _merged(_site_title(), gold_font, height=30, fill=navy_fill)
    _merged(report_title, Font(bold=True, size=13, color=NAVY), height=24)
    if subtitle:
        _merged(subtitle, muted_font, height=18)
    _merged(f"تاریخ تولید گزارش: {_now_jalali_str()}   —   تولیدکننده: {_generated_by()}", muted_font, height=18)
    row_cursor += 1  # ردیف خالی فاصله

    header_row = row_cursor
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=ci, value=h)
        c.font = white_font
        c.fill = navy_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[header_row].height = 22
    row_cursor += 1

    zebra_fill = PatternFill("solid", fgColor="F7F7F8")
    for ri, row in enumerate(rows):
        r = row_cursor + ri
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=ci, value=val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            cell.font = Font(size=10.5)
            if ri % 2 == 1:
                cell.fill = zebra_fill

    for ci, h in enumerate(headers, start=1):
        max_len = max([len(str(h))] + [len(str(r[ci - 1])) for r in rows if len(r) >= ci] or [0])
        ws.column_dimensions[get_column_letter(ci)].width = min(42, max(11, max_len + 4))

    ws.freeze_panes = f"A{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ─────────────────────────────────────────────────────────────────────────────
# PDF
# ─────────────────────────────────────────────────────────────────────────────

_FONT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "static", "fonts")
_FONTS_REGISTERED = False


def _register_fonts():
    """فونت وزیرمتن را یک‌بار برای کل عمر پروسه با reportlab ثبت می‌کند."""
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont("Vazir", os.path.join(_FONT_DIR, "Vazirmatn-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("Vazir-Bold", os.path.join(_FONT_DIR, "Vazirmatn-Bold.ttf")))
    _FONTS_REGISTERED = True


def _rtl(text) -> str:
    """شکل‌دهی حروف فارسی/عربی + ترتیب راست‌به‌چپ برای رسم صحیح در PDF."""
    import arabic_reshaper
    from bidi.algorithm import get_display
    s = "" if text is None else str(text)
    if not s:
        return s
    try:
        return get_display(arabic_reshaper.reshape(s))
    except Exception:
        return s


def export_pdf(headers, rows, report_title: str, filename: str, subtitle: str = "", landscape_mode: bool = None):
    """
    یک PDF تمیز و راست‌به‌چپ با سربرگ اطلاعاتی و جدول می‌سازد.
    اگر تعداد ستون‌ها زیاد باشد (>5)، خودکار landscape می‌شود مگر صریحاً مشخص شود.
    """
    try:
        _register_fonts()
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape as _landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
        )
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_CENTER
    except ImportError as e:
        # معمولاً یعنی بعد از آخرین بروزرسانی کد، `pip install -r requirements.txt`
        # روی سرور اجرا نشده و پکیج‌های reportlab/arabic-reshaper/python-bidi
        # نصب نیستند — پیام واضح برای عیب‌یابی سریع‌تر
        raise ExportDependencyError(
            f"وابستگی خروجی PDF نصب نیست ({e}). روی سرور دستور "
            "«pip install -r requirements.txt» را اجرا و سرویس را ری‌استارت کنید."
        ) from e

    if landscape_mode is None:
        landscape_mode = len(headers) > 5
    page_size = _landscape(A4) if landscape_mode else A4

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=page_size,
        rightMargin=14 * mm, leftMargin=14 * mm, topMargin=14 * mm, bottomMargin=16 * mm,
    )

    site_style = ParagraphStyle("site", fontName="Vazir-Bold", fontSize=15,
                                 alignment=TA_CENTER, textColor=colors.HexColor("#" + NAVY), spaceAfter=4)
    title_style = ParagraphStyle("title", fontName="Vazir-Bold", fontSize=12.5,
                                  alignment=TA_CENTER, textColor=colors.HexColor("#96703E"), spaceAfter=2)
    meta_style = ParagraphStyle("meta", fontName="Vazir", fontSize=9,
                                 alignment=TA_CENTER, textColor=colors.HexColor("#6B7280"), spaceAfter=2)

    elements = [
        Paragraph(_rtl(_site_title()), site_style),
        Paragraph(_rtl(report_title), title_style),
    ]
    if subtitle:
        elements.append(Paragraph(_rtl(subtitle), meta_style))
    elements.append(Paragraph(
        _rtl(f"تاریخ تولید گزارش: {_now_jalali_str()}   —   تولیدکننده: {_generated_by()}"), meta_style))
    elements.append(Spacer(1, 10 * mm))

    table_data = [[_rtl(h) for h in headers]]
    for row in rows:
        table_data.append([_rtl(v) for v in row])

    avail_width = page_size[0] - 28 * mm
    col_width = avail_width / max(1, len(headers))
    tbl = Table(table_data, colWidths=[col_width] * len(headers), repeatRows=1)
    tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Vazir-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Vazir"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E6E7EB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F8")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(tbl)

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Vazir", 8)
        canvas.setFillColor(colors.HexColor("#9CA3AF"))
        canvas.drawCentredString(page_size[0] / 2, 8 * mm,
                                  _rtl(f"صفحه {_doc.page} — سامانه دفترچه مالکیت {_site_title()}"))
        canvas.restoreState()

    doc.build(elements, onFirstPage=_footer, onLaterPages=_footer)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename, mimetype="application/pdf")
