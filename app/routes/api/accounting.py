"""API حسابداری تعاونی — حساب‌های دریافت وجوه و اسناد واریز

قواعد کسب‌وکار:
- سند واریز پس از ثبت قابل حذف نیست؛ فقط «ابطال» با ثبت دلیل (حسابرسی).
- شماره سند سریالی و یکتا است (از ۱۰۰۱ شروع می‌شود).
- عضو فقط اسناد واریز خودش را می‌بیند؛ ادمین همه را.
"""
import math
from datetime import datetime

from flask import Blueprint, jsonify, request
from flask_login import current_user, login_required
from sqlalchemy import func, or_

from ... import db_session
from ...models import (CoopAccount, DepositVoucher, ExpenseVoucher, MemberLedgerEntry,
                       Payment, Project, User, VoucherAllocation)
from ...utils.export import export_pdf, export_xlsx
from ...utils.files import FileValidationError, save_document
from ...utils.http import (get_request_body, has_finance_access, log_activity,
                           parse_int, require_admin, require_finance)
from ...utils.jalali import is_past_due
from ...utils.validation import normalize_digits

bp = Blueprint("accounting", __name__)

VOUCHER_METHODS  = ("card", "transfer", "cash", "cheque", "other")
EXPENSE_CATS     = ("land", "contractor", "admin", "utility", "other")


# ── کمک‌تابع‌های تخصیص/تسویه ─────────────────────────────────────────────────

def _payment_allocated(payment_id: int) -> int:
    """جمع تخصیص‌های فعال (از اسناد ابطال‌نشده) به یک قسط."""
    return int(db_session.query(func.coalesce(func.sum(VoucherAllocation.amount), 0))
               .join(DepositVoucher, VoucherAllocation.voucher_id == DepositVoucher.id)
               .filter(VoucherAllocation.payment_id == payment_id,
                       DepositVoucher.status == "active")
               .scalar())


def _recompute_payment_status(payment: Payment) -> None:
    """
    وضعیت قسط را بر اساس مجموع تخصیص‌ها بازمحاسبه می‌کند:
    - تخصیص کامل → paid
    - ناقص و سررسید گذشته → overdue
    - ناقص و سررسید نگذشته → unpaid
    """
    allocated = _payment_allocated(payment.id)
    if allocated >= payment.amount:
        payment.status = "paid"
        if not payment.paid_date:
            from ...utils.jalali import today_jalali
            payment.paid_date = today_jalali()
    else:
        payment.paid_date = None
        payment.status = "overdue" if is_past_due(payment.due_date) else "unpaid"


def _voucher_used(voucher_id: int) -> int:
    """جمع مبالغ تخصیص‌یافته از یک سند واریز."""
    return int(db_session.query(func.coalesce(func.sum(VoucherAllocation.amount), 0))
               .filter(VoucherAllocation.voucher_id == voucher_id).scalar())


# ── حساب‌های تعاونی ───────────────────────────────────────────────────────────

@bp.get("/accounting/accounts")
@login_required
def accounts_list():
    err = require_finance()
    if err: return err
    items = db_session.query(CoopAccount).order_by(CoopAccount.created_at).all()
    return jsonify(accounts=[a.to_dict() for a in items])


@bp.post("/accounting/accounts")
@login_required
def accounts_create():
    err = require_finance()
    if err: return err

    data = get_request_body()
    holder = (data.get("holder_name") or "").strip()
    if not holder:
        return jsonify(error="نام صاحب حساب الزامی است."), 400

    a = CoopAccount(
        holder_name=holder,
        holder_role=(data.get("holder_role") or "").strip(),
        bank_name=(data.get("bank_name") or "").strip(),
        account_number=normalize_digits((data.get("account_number") or "").strip()),
        card_number=normalize_digits((data.get("card_number") or "").strip()),
        iban=normalize_digits((data.get("iban") or "").strip()),
    )
    db_session.add(a)
    db_session.commit()
    log_activity("account_create", f"حساب جدید: {holder}", "finance")
    return jsonify(ok=True, account=a.to_dict()), 201


@bp.post("/accounting/accounts/<int:aid>")
@login_required
def accounts_update(aid):
    err = require_finance()
    if err: return err

    a = db_session.get(CoopAccount, aid)
    if not a:
        return jsonify(error="حساب یافت نشد."), 404

    data = get_request_body()
    for field in ("holder_name", "holder_role", "bank_name"):
        if data.get(field) is not None:
            setattr(a, field, str(data[field]).strip())
    for field in ("account_number", "card_number", "iban"):
        if data.get(field) is not None:
            setattr(a, field, normalize_digits(str(data[field]).strip()))
    if "active" in data:
        a.active = bool(data["active"])

    db_session.commit()
    log_activity("account_update", f"ویرایش حساب: {a.holder_name}", "finance")
    return jsonify(ok=True)


# ── اسناد واریز ───────────────────────────────────────────────────────────────

@bp.get("/accounting/vouchers")
@login_required
def vouchers_list():
    """
    لیست اسناد با فیلتر و صفحه‌بندی.
    فیلترها: account_id، user_id، status، date_from/date_to (شمسی)، search.
    خروجی شامل جمع مبالغ اسناد فعالِ منطبق با فیلتر است.
    """
    is_finance = has_finance_access()
    q = db_session.query(DepositVoucher)
    if not is_finance:
        q = q.filter(DepositVoucher.user_id == current_user.id)

    account_id = request.args.get("account_id", type=int)
    user_id    = request.args.get("user_id", type=int)
    status     = (request.args.get("status") or "").strip()
    date_from  = (request.args.get("date_from") or "").strip()
    date_to    = (request.args.get("date_to") or "").strip()
    search     = (request.args.get("search") or "").strip()
    unmatched  = request.args.get("unmatched") == "1"
    page       = request.args.get("page", 1, type=int)
    per_page   = min(request.args.get("per_page", 30, type=int) or 30, 200)

    if account_id: q = q.filter(DepositVoucher.account_id == account_id)
    if user_id and is_finance:
        q = q.filter(DepositVoucher.user_id == user_id)
    if status:     q = q.filter(DepositVoucher.status == status)
    if unmatched and is_finance:
        q = q.filter(DepositVoucher.user_id.is_(None), DepositVoucher.source == "bank_import")
    if date_from:  q = q.filter(DepositVoucher.deposit_date >= date_from)
    if date_to:    q = q.filter(DepositVoucher.deposit_date <= date_to)
    if search:
        like = f"%{search}%"
        q = (q.outerjoin(User, DepositVoucher.user_id == User.id)
              .filter(or_(
                  User.first_name.ilike(like), User.last_name.ilike(like),
                  DepositVoucher.payer_name.ilike(like),
                  DepositVoucher.reference_no.ilike(like),
                  DepositVoucher.description.ilike(like),
              )))

    total = q.count()
    total_amount = (q.filter(DepositVoucher.status == "active")
                    .with_entities(func.coalesce(func.sum(DepositVoucher.amount), 0))
                    .scalar())

    items = (q.order_by(DepositVoucher.number.desc())
             .offset((page - 1) * per_page).limit(per_page).all())

    return jsonify(
        vouchers=[v.to_dict() for v in items],
        total=total, total_amount=int(total_amount),
        page=page, pages=math.ceil(total / per_page) or 1,
    )


@bp.post("/accounting/vouchers")
@login_required
def vouchers_create():
    """ثبت سند واریز جدید (فقط ادمین) — multipart برای فیش یا JSON"""
    err = require_finance()
    if err: return err

    data = get_request_body()

    account = db_session.get(CoopAccount, parse_int(data.get("account_id")))
    if not account:
        return jsonify(error="حساب مقصد را انتخاب کنید."), 400

    amount = parse_int(normalize_digits(str(data.get("amount", "")).replace(",", "")))
    if amount <= 0:
        return jsonify(error="مبلغ واریز نامعتبر است."), 400

    deposit_date = normalize_digits((data.get("deposit_date") or "").strip())
    if not deposit_date:
        return jsonify(error="تاریخ واریز الزامی است."), 400

    user_id    = parse_int(data.get("user_id")) or None
    payer_name = (data.get("payer_name") or "").strip()
    if user_id:
        payer = db_session.get(User, user_id)
        if not payer:
            return jsonify(error="عضو واریزکننده یافت نشد."), 404
    elif not payer_name:
        return jsonify(error="واریزکننده را مشخص کنید (عضو یا نام آزاد)."), 400

    method = data.get("method", "card")
    if method not in VOUCHER_METHODS:
        method = "other"

    receipt_path = None
    f = request.files.get("receipt")
    if f and f.filename:
        try:
            receipt_path = save_document(f, "payments")
        except FileValidationError as e:
            return jsonify(error=str(e)), 400

    last_no = db_session.query(func.max(DepositVoucher.number)).scalar() or 1000
    v = DepositVoucher(
        number=last_no + 1,
        user_id=user_id,
        payer_name=payer_name if not user_id else "",
        account_id=account.id,
        amount=amount,
        deposit_date=deposit_date,
        deposit_time=normalize_digits((data.get("deposit_time") or "").strip()),
        method=method,
        reference_no=normalize_digits((data.get("reference_no") or "").strip()),
        description=data.get("description") or None,
        receipt_path=receipt_path,
        payment_id=parse_int(data.get("payment_id")) or None,
        created_by=current_user.id,
    )
    db_session.add(v)
    db_session.commit()
    log_activity("voucher_create",
                 f"سند #{v.number}: {v.to_dict()['payer_name']} — {amount:,} ریال به حساب {account.holder_name}",
                 "finance")
    return jsonify(ok=True, voucher=v.to_dict()), 201


@bp.post("/accounting/vouchers/<int:vid>/void")
@login_required
def vouchers_void(vid):
    """ابطال سند — حذف فیزیکی مجاز نیست"""
    err = require_finance()
    if err: return err

    v = db_session.get(DepositVoucher, vid)
    if not v:
        return jsonify(error="سند یافت نشد."), 404
    if v.status == "voided":
        return jsonify(error="سند قبلاً ابطال شده."), 400

    reason = (get_request_body().get("reason") or "").strip()
    if not reason:
        return jsonify(error="دلیل ابطال الزامی است."), 400

    v.status      = "voided"
    v.void_reason = reason
    v.voided_at   = datetime.utcnow()

    # تخصیص‌های این سند حذف و اقساط متأثر بازمحاسبه می‌شوند
    affected = db_session.query(VoucherAllocation).filter_by(voucher_id=v.id).all()
    payment_ids = {a.payment_id for a in affected}
    for a in affected:
        db_session.delete(a)
    db_session.flush()
    for pid in payment_ids:
        p = db_session.get(Payment, pid)
        if p:
            _recompute_payment_status(p)

    db_session.commit()
    log_activity("voucher_void", f"ابطال سند #{v.number} — {reason}", "finance")
    return jsonify(ok=True)


@bp.post("/accounting/vouchers/import-bank-statement")
@login_required
def vouchers_import_bank_statement():
    """
    وارد کردن صورت‌حساب خام بانکی (خروجی مستقیم سامانه بانکداری) و تطبیق خودکار
    هر ردیف واریزی با عضوی که «شناسه واریز» او داخل توضیحات تراکنش پیدا شود.
    ردیف‌های بدون تطبیق هم به‌عنوان سند (بدون عضو) ثبت می‌شوند تا مبلغ در حساب دیده شود،
    و بعداً از فهرست «بدون تطبیق» قابل اصلاح‌اند (ابطال و ثبت مجدد با عضو درست).
    """
    err = require_finance()
    if err: return err

    account = db_session.get(CoopAccount, parse_int(request.form.get("account_id")))
    if not account:
        return jsonify(error="حساب مقصد را انتخاب کنید."), 400

    if "file" not in request.files:
        return jsonify(error="فایل انتخاب نشده."), 400
    f = request.files["file"]
    if not f.filename.endswith((".xlsx", ".xls")):
        return jsonify(error="فرمت فایل باید xlsx باشد."), 400

    from openpyxl import load_workbook

    from ...utils.bank_import import (build_deposit_id_map, cell_amount, cell_str,
                                       extract_deposit_id_candidates, find_bank_header)

    try:
        wb = load_workbook(filename=f, data_only=True)
    except Exception:
        return jsonify(error="فایل اکسل معتبر نیست."), 400
    ws = wb.active

    header_row, col_idx = find_bank_header(ws)
    if header_row is None:
        return jsonify(error="ستون‌های صورت‌حساب پیدا نشد — مطمئن شوید فایل خروجی مستقیم بانک است."), 400

    dep_map = build_deposit_id_map(
        db_session.query(User).filter(User.deposit_id.isnot(None), User.deposit_id != "").all()
    )

    def g(row, key):
        i = col_idx.get(key)
        return row[i] if i is not None and i < len(row) else None

    last_no = db_session.query(func.max(DepositVoucher.number)).scalar() or 1000
    matched = unmatched = dup_count = zero_count = 0

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        # صورت‌حساب‌های بانکی گاهی بعد از آخرین تراکنش، چند ردیف خالی و بعد یک
        # جدول خلاصه/گزارش کلی (با اعداد نامرتبط در همان ستون‌ها) دارند —
        # تراکنش واقعی هرگز ردیف کاملاً خالی میانش ندارد، پس با رسیدن به اولین
        # ردیف خالی متوقف می‌شویم، نه فقط رد کردنش، وگرنه ردیف‌های خلاصه‌ی بعدی
        # به‌عنوان واریزی ساختگی با مبالغ نجومی وارد می‌شوند.
        if not row or not any(row):
            break

        amount = cell_amount(g(row, "deposit_amount"))
        if amount <= 0:
            zero_count += 1
            continue

        deposit_date = normalize_digits(cell_str(g(row, "date")).strip())
        deposit_time = normalize_digits(cell_str(g(row, "time")).strip())
        reference_no = normalize_digits(cell_str(g(row, "reference_no")))
        branch_code  = cell_str(g(row, "branch_code"))
        branch_name  = cell_str(g(row, "branch_name"))
        channel      = cell_str(g(row, "channel"))
        balance      = cell_amount(g(row, "balance")) or None
        desc_raw     = " ".join(filter(None, (
            cell_str(row[i]) for i in col_idx["desc_cols"] if i < len(row)
        )))

        dup = db_session.query(DepositVoucher).filter_by(
            account_id=account.id, reference_no=reference_no,
            deposit_date=deposit_date, deposit_time=deposit_time, amount=amount,
        ).first()
        if dup:
            dup_count += 1
            continue

        user = None
        for cand in extract_deposit_id_candidates(desc_raw):
            user = dep_map.get(cand)
            if user:
                break

        last_no += 1
        v = DepositVoucher(
            number=last_no,
            user_id=user.id if user else None,
            payer_name="" if user else (desc_raw[:120] or "نامشخص"),
            account_id=account.id,
            amount=amount,
            deposit_date=deposit_date,
            deposit_time=deposit_time,
            method="cash" if "نقد" in desc_raw else "transfer",
            reference_no=reference_no,
            description=desc_raw or None,
            source="bank_import",
            branch_code=branch_code,
            branch_name=branch_name,
            channel=channel,
            bank_balance_after=balance,
            created_by=current_user.id,
        )
        db_session.add(v)
        if user:
            matched += 1
        else:
            unmatched += 1

    db_session.commit()
    log_activity(
        "voucher_bank_import",
        f"واردات صورت‌حساب بانکی حساب {account.holder_name}: "
        f"{matched} واریزی تطبیق‌یافته، {unmatched} بدون تطبیق، {dup_count} تکراری نادیده‌گرفته‌شده",
        "finance",
    )
    return jsonify(ok=True, matched=matched, unmatched=unmatched,
                   skipped_duplicate=dup_count, skipped_zero=zero_count), 201


@bp.get("/accounting/summary")
@login_required
def accounting_summary():
    """موجودی هر حساب = واریزی‌های فعال − هزینه‌های فعال."""
    err = require_finance()
    if err: return err

    result = []
    for a in db_session.query(CoopAccount).order_by(CoopAccount.created_at).all():
        deposits = int(db_session.query(func.coalesce(func.sum(DepositVoucher.amount), 0))
                       .filter_by(account_id=a.id, status="active").scalar())
        d_count  = db_session.query(func.count(DepositVoucher.id)).filter_by(
                       account_id=a.id, status="active").scalar()
        expenses = int(db_session.query(func.coalesce(func.sum(ExpenseVoucher.amount), 0))
                       .filter_by(account_id=a.id, status="active").scalar())
        e_count  = db_session.query(func.count(ExpenseVoucher.id)).filter_by(
                       account_id=a.id, status="active").scalar()
        # آخرین «مانده» گزارش‌شده در صورت‌حساب بانکی وارد‌شده — برای تطبیق دستی
        # موجودی محاسبه‌شده سامانه با آخرین مانده اعلامی خود بانک
        last_import = (db_session.query(DepositVoucher)
                       .filter(DepositVoucher.account_id == a.id,
                               DepositVoucher.source == "bank_import",
                               DepositVoucher.bank_balance_after.isnot(None))
                       .order_by(DepositVoucher.deposit_date.desc(),
                                 DepositVoucher.deposit_time.desc(),
                                 DepositVoucher.id.desc())
                       .first())
        result.append({
            "account": a.to_dict(),
            "total_amount": deposits,          # سازگاری قبلی
            "deposits": deposits, "expenses": expenses,
            "balance": deposits - expenses,
            "voucher_count": d_count, "expense_count": e_count,
            "bank_reported_balance": last_import.bank_balance_after if last_import else None,
            "bank_reported_at": (f"{last_import.deposit_date} {last_import.deposit_time}".strip()
                                 if last_import else None),
        })
    return jsonify(summary=result)


# ── تخصیص سند به اقساط (تسویه) ───────────────────────────────────────────────

@bp.get("/accounting/vouchers/<int:vid>/allocations")
@login_required
def voucher_allocations(vid):
    """تخصیص‌های فعلی یک سند + اقساط بازِ عضوِ واریزکننده برای تخصیص."""
    err = require_finance()
    if err: return err

    v = db_session.get(DepositVoucher, vid)
    if not v:
        return jsonify(error="سند یافت نشد."), 404

    allocs = db_session.query(VoucherAllocation).filter_by(voucher_id=vid).all()
    open_payments = []
    if v.user_id:
        for p in (db_session.query(Payment)
                  .filter(Payment.user_id == v.user_id)
                  .order_by(Payment.due_date).all()):
            allocated = _payment_allocated(p.id)
            open_payments.append({
                "id": p.id, "description": p.description or "",
                "due_date": p.due_date, "amount": p.amount,
                "allocated": allocated, "remaining": max(0, p.amount - allocated),
                "status": p.status,
            })

    return jsonify(
        voucher=v.to_dict(),
        used=_voucher_used(vid),
        remaining=max(0, v.amount - _voucher_used(vid)),
        allocations=[a.to_dict() for a in allocs],
        open_payments=open_payments,
    )


@bp.post("/accounting/vouchers/<int:vid>/allocate")
@login_required
def voucher_allocate(vid):
    """
    تخصیص سند به اقساط. بدنه: {"items": [{"payment_id": .., "amount": ..}, ...]}
    تخصیص‌های قبلی این سند جایگزین می‌شوند و اقساط متأثر بازمحاسبه می‌شوند.
    """
    err = require_finance()
    if err: return err

    v = db_session.get(DepositVoucher, vid)
    if not v:
        return jsonify(error="سند یافت نشد."), 404
    if v.status != "active":
        return jsonify(error="سند ابطال‌شده قابل تخصیص نیست."), 400

    items = get_request_body().get("items") or []
    cleaned, seen = [], set()
    total = 0
    for it in items:
        pid = parse_int(it.get("payment_id"))
        amt = parse_int(it.get("amount"))
        if not pid or amt <= 0:
            continue
        if pid in seen:
            return jsonify(error="قسط تکراری در تخصیص."), 400
        seen.add(pid)
        p = db_session.get(Payment, pid)
        if not p:
            return jsonify(error=f"قسط {pid} یافت نشد."), 404
        if v.user_id and p.user_id != v.user_id:
            return jsonify(error="قسط متعلق به واریزکننده نیست."), 400
        cleaned.append((p, amt))
        total += amt

    if total > v.amount:
        return jsonify(error=f"جمع تخصیص ({total:,}) از مبلغ سند ({v.amount:,}) بیشتر است."), 400

    # جایگزینی تخصیص‌های قبلی
    prev = db_session.query(VoucherAllocation).filter_by(voucher_id=vid).all()
    prev_pids = {a.payment_id for a in prev}
    for a in prev:
        db_session.delete(a)
    db_session.flush()

    for p, amt in cleaned:
        db_session.add(VoucherAllocation(voucher_id=vid, payment_id=p.id, amount=amt))
    db_session.flush()

    # بازمحاسبه وضعیت همه اقساط متأثر (قبلی + جدید)
    for pid in prev_pids | {p.id for p, _ in cleaned}:
        p = db_session.get(Payment, pid)
        if p:
            _recompute_payment_status(p)

    db_session.commit()
    log_activity("voucher_allocate",
                 f"تخصیص سند #{v.number} به {len(cleaned)} قسط ({total:,} ریال)", "finance")
    return jsonify(ok=True, used=total, remaining=max(0, v.amount - total))


# ── کارت حساب عضو (کاردکس) ────────────────────────────────────────────────────

def _build_statement(u: User) -> dict:
    """داده صورت‌حساب یک عضو را می‌سازد (مشترک بین JSON و Excel)."""
    payments = (db_session.query(Payment).filter_by(user_id=u.id)
                .order_by(Payment.due_date).all())
    vouchers = (db_session.query(DepositVoucher)
                .filter_by(user_id=u.id, status="active")
                .order_by(DepositVoucher.deposit_date).all())
    ledger_entries = (db_session.query(MemberLedgerEntry)
                      .filter_by(user_id=u.id, status="active")
                      .order_by(MemberLedgerEntry.entry_date).all())

    manual_debit  = sum(e.amount for e in ledger_entries if e.entry_type == "debit")
    manual_credit = sum(e.amount for e in ledger_entries if e.entry_type == "credit")

    total_debt      = sum(p.amount for p in payments) + manual_debit
    total_paid      = sum(_payment_allocated(p.id) for p in payments) + manual_credit
    total_deposited = sum(v.amount for v in vouchers)
    unallocated     = total_deposited - sum(_voucher_used(v.id) for v in vouchers)

    tx = []
    for p in payments:
        tx.append({"source": "payment", "id": p.id, "type": "debit", "date": p.due_date,
                   "title": p.description or "قسط", "amount": p.amount,
                   "allocated": _payment_allocated(p.id), "status": p.status})
    for v in vouchers:
        tx.append({"source": "voucher", "id": v.id, "type": "credit", "date": v.deposit_date,
                   "title": f"واریز — سند #{v.number}", "amount": v.amount,
                   "reference_no": v.reference_no or ""})
    for e in ledger_entries:
        tx.append({"source": "ledger", "id": e.id, "type": e.entry_type, "date": e.entry_date,
                   "title": e.title, "amount": e.amount, "note": e.note or ""})
    tx.sort(key=lambda t: t["date"])

    # مانده تجمعی (کاردکس) — بعد از هر تراکنش، به ترتیب تاریخ
    balance = 0
    for t in tx:
        balance += t["amount"] if t["type"] == "debit" else -t["amount"]
        t["balance"] = balance

    return {
        "member": {"id": u.id, "full_name": u.full_name,
                   "national_code": u.national_code or ""},
        "totals": {"debt": total_debt, "paid": total_paid,
                   "remaining_debt": max(0, total_debt - total_paid),
                   "deposited": total_deposited,
                   "unallocated": max(0, unallocated),
                   "balance": balance},
        "transactions": tx,
    }


@bp.get("/accounting/member/<int:uid>/statement")
@login_required
def member_statement(uid):
    """صورت‌حساب عضو: بدهی، واریزی، مانده و ریز تراکنش‌ها."""
    if not has_finance_access() and current_user.id != uid:
        return jsonify(error="دسترسی ندارید."), 403
    u = db_session.get(User, uid)
    if not u:
        return jsonify(error="عضو یافت نشد."), 404
    return jsonify(**_build_statement(u))


@bp.get("/accounting/member/<int:uid>/statement/export")
@login_required
def member_statement_export(uid):
    """خروجی Excel کارت حساب عضو."""
    if not has_finance_access() and current_user.id != uid:
        return jsonify(error="دسترسی ندارید."), 403
    u = db_session.get(User, uid)
    if not u:
        return jsonify(error="عضو یافت نشد."), 404

    st = _build_statement(u)
    rows = []
    for t in st["transactions"]:
        rows.append([
            t["date"],
            t["title"],
            t["amount"] if t["type"] == "debit" else "",     # بدهکار
            t["amount"] if t["type"] == "credit" else "",    # بستانکار
            ("پرداخت‌شده" if t.get("status") == "paid"
             else "معوق" if t.get("status") == "overdue"
             else "") if t["source"] == "payment" else (t.get("reference_no") or t.get("note") or ""),
            t["balance"],
        ])
    tot = st["totals"]
    rows.append([])
    rows.append(["جمع", "", tot["debt"], tot["paid"], "", ""])
    rows.append(["مانده بدهی", "", tot["remaining_debt"], "", "", ""])

    headers = ["تاریخ", "شرح", "بدهکار", "بستانکار", "وضعیت/پیگیری", "مانده"]
    title = f"کارت حساب {u.full_name}"
    fmt = (request.args.get("format") or "xlsx").strip()
    if fmt == "pdf":
        return export_pdf(headers, rows, title, f"statement_{uid}.pdf")
    return export_xlsx(headers, rows, title, f"statement_{uid}.xlsx")


@bp.post("/accounting/member/<int:uid>/ledger")
@login_required
def member_ledger_create(uid):
    """ثبت سند دستی بدهکار/بستانکار در کارت حساب عضو (جریمه، تخفیف، اصلاح حساب و ...)."""
    err = require_finance()
    if err: return err

    u = db_session.get(User, uid)
    if not u:
        return jsonify(error="عضو یافت نشد."), 404

    data       = get_request_body()
    entry_type = data.get("entry_type")
    title      = (data.get("title") or "").strip()
    entry_date = (data.get("entry_date") or "").strip()
    amount     = parse_int(data.get("amount"))

    if entry_type not in ("debit", "credit"):
        return jsonify(error="نوع سند باید بدهکار یا بستانکار باشد."), 400
    if not title or not entry_date or not amount or amount <= 0:
        return jsonify(error="عنوان، تاریخ و مبلغ (بزرگ‌تر از صفر) الزامی است."), 400

    entry = MemberLedgerEntry(
        user_id=uid, entry_type=entry_type, amount=amount,
        entry_date=entry_date, title=title, note=data.get("note", ""),
        created_by=current_user.id,
    )
    db_session.add(entry)
    db_session.commit()
    log_activity("ledger_create",
                 f"سند {'بدهکار' if entry_type=='debit' else 'بستانکار'} برای {u.full_name}: "
                 f"{title} — {amount:,} ریال", "finance")
    return jsonify(ok=True, entry=entry.to_dict()), 201


@bp.post("/accounting/ledger/<int:eid>/void")
@login_required
def member_ledger_void(eid):
    """ابطال سند دستی — رکورد برای حسابرسی باقی می‌ماند، فقط از محاسبات خارج می‌شود."""
    err = require_finance()
    if err: return err

    entry = db_session.get(MemberLedgerEntry, eid)
    if not entry or entry.status == "voided":
        return jsonify(error="یافت نشد."), 404

    data = get_request_body()
    entry.status      = "voided"
    entry.void_reason = data.get("reason", "")
    entry.voided_at   = datetime.utcnow()
    db_session.commit()
    log_activity("ledger_void", f"ابطال سند دستی #{entry.id} — {entry.title}", "finance")
    return jsonify(ok=True)


# ── اسناد هزینه (خروجی وجه) ──────────────────────────────────────────────────

@bp.get("/accounting/expenses")
@login_required
def expenses_list():
    err = require_finance()
    if err: return err

    q = db_session.query(ExpenseVoucher)
    account_id = request.args.get("account_id", type=int)
    project_id = request.args.get("project_id", type=int)
    category   = (request.args.get("category") or "").strip()
    status     = (request.args.get("status") or "").strip()
    search     = (request.args.get("search") or "").strip()
    page       = request.args.get("page", 1, type=int)
    per_page   = min(request.args.get("per_page", 30, type=int) or 30, 200)

    if account_id: q = q.filter(ExpenseVoucher.account_id == account_id)
    if project_id: q = q.filter(ExpenseVoucher.project_id == project_id)
    if category:   q = q.filter(ExpenseVoucher.category == category)
    if status:     q = q.filter(ExpenseVoucher.status == status)
    if search:
        like = f"%{search}%"
        q = q.filter(or_(ExpenseVoucher.payee.ilike(like),
                         ExpenseVoucher.reference_no.ilike(like),
                         ExpenseVoucher.description.ilike(like)))

    total = q.count()
    total_amount = int(q.filter(ExpenseVoucher.status == "active")
                       .with_entities(func.coalesce(func.sum(ExpenseVoucher.amount), 0)).scalar())
    items = (q.order_by(ExpenseVoucher.number.desc())
             .offset((page - 1) * per_page).limit(per_page).all())
    return jsonify(expenses=[e.to_dict() for e in items], total=total,
                   total_amount=total_amount, page=page,
                   pages=math.ceil(total / per_page) or 1)


@bp.post("/accounting/expenses")
@login_required
def expenses_create():
    err = require_finance()
    if err: return err

    data = get_request_body()
    account = db_session.get(CoopAccount, parse_int(data.get("account_id")))
    if not account:
        return jsonify(error="حساب پرداخت‌کننده را انتخاب کنید."), 400

    amount = parse_int(normalize_digits(str(data.get("amount", "")).replace(",", "")))
    if amount <= 0:
        return jsonify(error="مبلغ هزینه نامعتبر است."), 400

    spend_date = normalize_digits((data.get("spend_date") or "").strip())
    if not spend_date:
        return jsonify(error="تاریخ هزینه الزامی است."), 400

    payee = (data.get("payee") or "").strip()
    if not payee:
        return jsonify(error="دریافت‌کننده (بابت) را مشخص کنید."), 400

    category = data.get("category", "other")
    if category not in EXPENSE_CATS:
        category = "other"

    receipt_path = None
    f = request.files.get("receipt")
    if f and f.filename:
        try:
            receipt_path = save_document(f, "payments")
        except FileValidationError as e:
            return jsonify(error=str(e)), 400

    last_no = db_session.query(func.max(ExpenseVoucher.number)).scalar() or 5000
    e = ExpenseVoucher(
        number=last_no + 1,
        account_id=account.id,
        project_id=parse_int(data.get("project_id")) or None,
        category=category, payee=payee, amount=amount,
        spend_date=spend_date,
        spend_time=normalize_digits((data.get("spend_time") or "").strip()),
        method=data.get("method", "transfer"),
        reference_no=normalize_digits((data.get("reference_no") or "").strip()),
        description=data.get("description") or None,
        receipt_path=receipt_path,
        created_by=current_user.id,
    )
    db_session.add(e)
    db_session.commit()
    log_activity("expense_create",
                 f"هزینه #{e.number}: {payee} — {amount:,} ریال از حساب {account.holder_name}", "finance")
    return jsonify(ok=True, expense=e.to_dict()), 201


@bp.post("/accounting/expenses/<int:eid>/void")
@login_required
def expenses_void(eid):
    err = require_finance()
    if err: return err

    e = db_session.get(ExpenseVoucher, eid)
    if not e:
        return jsonify(error="سند یافت نشد."), 404
    if e.status == "voided":
        return jsonify(error="سند قبلاً ابطال شده."), 400

    reason = (get_request_body().get("reason") or "").strip()
    if not reason:
        return jsonify(error="دلیل ابطال الزامی است."), 400

    e.status      = "voided"
    e.void_reason = reason
    e.voided_at   = datetime.utcnow()
    db_session.commit()
    log_activity("expense_void", f"ابطال هزینه #{e.number} — {reason}", "finance")
    return jsonify(ok=True)


# ── داشبورد مالی و تفکیک پروژه ───────────────────────────────────────────────

@bp.get("/accounting/dashboard")
@login_required
def accounting_dashboard():
    err = require_finance()
    if err: return err

    total_deposited = int(db_session.query(func.coalesce(func.sum(DepositVoucher.amount), 0))
                          .filter_by(status="active").scalar())
    total_expenses  = int(db_session.query(func.coalesce(func.sum(ExpenseVoucher.amount), 0))
                          .filter_by(status="active").scalar())

    # اقساط
    total_billed = int(db_session.query(func.coalesce(func.sum(Payment.amount), 0)).scalar())
    total_collected = int(db_session.query(func.coalesce(func.sum(VoucherAllocation.amount), 0))
                          .join(DepositVoucher, VoucherAllocation.voucher_id == DepositVoucher.id)
                          .filter(DepositVoucher.status == "active").scalar())
    overdue_q = db_session.query(Payment).filter(Payment.status == "overdue")
    overdue_count  = overdue_q.count()
    overdue_amount = int(overdue_q.with_entities(
        func.coalesce(func.sum(Payment.amount), 0)).scalar())

    return jsonify(
        total_deposited=total_deposited,
        total_expenses=total_expenses,
        balance=total_deposited - total_expenses,
        total_billed=total_billed,
        total_collected=total_collected,
        collection_rate=round(total_collected / total_billed * 100, 1) if total_billed else 0,
        overdue_count=overdue_count,
        overdue_amount=overdue_amount,
    )


@bp.get("/accounting/projects-summary")
@login_required
def projects_financial_summary():
    """تفکیک مالی به‌ازای هر پروژه: بدهی، وصولی، هزینه، درصد وصول."""
    err = require_finance()
    if err: return err

    from ...models import Booklet
    result = []
    for proj in db_session.query(Project).order_by(Project.order, Project.name).all():
        billed = int(db_session.query(func.coalesce(func.sum(Payment.amount), 0))
                     .join(Booklet, Payment.booklet_id == Booklet.id)
                     .filter(Booklet.project_id == proj.id).scalar())
        collected = int(db_session.query(func.coalesce(func.sum(VoucherAllocation.amount), 0))
                        .join(DepositVoucher, VoucherAllocation.voucher_id == DepositVoucher.id)
                        .join(Payment, VoucherAllocation.payment_id == Payment.id)
                        .join(Booklet, Payment.booklet_id == Booklet.id)
                        .filter(Booklet.project_id == proj.id,
                                DepositVoucher.status == "active").scalar())
        expenses = int(db_session.query(func.coalesce(func.sum(ExpenseVoucher.amount), 0))
                       .filter_by(project_id=proj.id, status="active").scalar())
        result.append({
            "project_id": proj.id, "project_name": proj.name,
            "billed": billed, "collected": collected,
            "remaining": max(0, billed - collected),
            "expenses": expenses,
            "collection_rate": round(collected / billed * 100, 1) if billed else 0,
        })
    return jsonify(projects=result)


@bp.get("/accounting/members-balances")
@login_required
def members_balances():
    """
    فهرست بدهکار/بستانکار همه اعضا برای پیگیری مالی — یک نمای کلی سریع
    به‌جای باز کردن کارت حساب تک‌تک اعضا.
    """
    err = require_finance()
    if err: return err

    debt_rows = dict(db_session.query(Payment.user_id, func.coalesce(func.sum(Payment.amount), 0))
                     .group_by(Payment.user_id).all())
    paid_rows = dict(
        db_session.query(Payment.user_id, func.coalesce(func.sum(VoucherAllocation.amount), 0))
        .join(VoucherAllocation, VoucherAllocation.payment_id == Payment.id)
        .join(DepositVoucher, VoucherAllocation.voucher_id == DepositVoucher.id)
        .filter(DepositVoucher.status == "active")
        .group_by(Payment.user_id).all())
    deposited_rows = dict(
        db_session.query(DepositVoucher.user_id, func.coalesce(func.sum(DepositVoucher.amount), 0))
        .filter(DepositVoucher.status == "active", DepositVoucher.user_id.isnot(None))
        .group_by(DepositVoucher.user_id).all())
    ledger_debit = dict(
        db_session.query(MemberLedgerEntry.user_id, func.coalesce(func.sum(MemberLedgerEntry.amount), 0))
        .filter(MemberLedgerEntry.status == "active", MemberLedgerEntry.entry_type == "debit")
        .group_by(MemberLedgerEntry.user_id).all())
    ledger_credit = dict(
        db_session.query(MemberLedgerEntry.user_id, func.coalesce(func.sum(MemberLedgerEntry.amount), 0))
        .filter(MemberLedgerEntry.status == "active", MemberLedgerEntry.entry_type == "credit")
        .group_by(MemberLedgerEntry.user_id).all())

    result = []
    for m in db_session.query(User).filter_by(role="member").all():
        debt      = int(debt_rows.get(m.id, 0)) + int(ledger_debit.get(m.id, 0))
        paid      = int(paid_rows.get(m.id, 0)) + int(ledger_credit.get(m.id, 0))
        deposited = int(deposited_rows.get(m.id, 0))
        if not (debt or paid or deposited):
            continue
        result.append({
            "id": m.id, "full_name": m.full_name, "national_code": m.national_code or "",
            "personnel_code": m.personnel_code or "", "org_unit": m.org_unit or "",
            "debt": debt, "paid": paid, "remaining_debt": max(0, debt - paid),
            "deposited": deposited,
        })
    result.sort(key=lambda x: -x["remaining_debt"])
    return jsonify(members=result)


@bp.get("/accounting/org-units-summary")
@login_required
def org_units_summary():
    """
    جمع واریزی‌های تطبیق‌یافته (متعلق به یک عضو) به تفکیک «واحد سازمانی» —
    برای مقایسه میزان مشارکت هر واحد/استان در تأمین مالی پروژه.
    ردیف‌های بدون واحد سازمانی ثبت‌شده زیر عنوان «نامشخص» جمع می‌شوند.
    """
    err = require_finance()
    if err: return err

    rows = (db_session.query(func.coalesce(User.org_unit, "نامشخص"),
                             func.coalesce(func.sum(DepositVoucher.amount), 0),
                             func.count(DepositVoucher.id),
                             func.count(func.distinct(User.id)))
            .join(User, DepositVoucher.user_id == User.id)
            .filter(DepositVoucher.status == "active")
            .group_by(func.coalesce(User.org_unit, "نامشخص"))
            .order_by(func.sum(DepositVoucher.amount).desc())
            .all())

    result = [
        {"org_unit": org_unit, "total_amount": int(total), "voucher_count": int(cnt), "member_count": int(mcnt)}
        for org_unit, total, cnt, mcnt in rows
    ]
    grand_total = sum(r["total_amount"] for r in result)
    return jsonify(org_units=result, grand_total=grand_total)


# ── خروجی Excel/PDF ─────────────────────────────────────────────────────────

@bp.get("/accounting/vouchers/export")
@login_required
def vouchers_export():
    err = require_finance()
    if err: return err
    q = db_session.query(DepositVoucher)
    account_id = request.args.get("account_id", type=int)
    status     = (request.args.get("status") or "").strip()
    if account_id: q = q.filter(DepositVoucher.account_id == account_id)
    if status:     q = q.filter(DepositVoucher.status == status)
    vs = q.order_by(DepositVoucher.number).all()
    method_lbl = {"card": "کارت", "transfer": "انتقال", "cash": "نقدی", "cheque": "چک", "other": "سایر"}
    rows = [[
        v.number, v.to_dict()["payer_name"],
        v.account.holder_name if v.account else "",
        v.amount, v.deposit_date, v.deposit_time or "",
        method_lbl.get(v.method, v.method), v.reference_no or "",
        "ابطال" if v.status == "voided" else "فعال",
        v.description or "",
    ] for v in vs]
    headers = ["شماره سند", "واریزکننده", "حساب مقصد", "مبلغ (ریال)", "تاریخ", "ساعت",
               "روش", "شماره پیگیری", "وضعیت", "شرح"]
    fmt = (request.args.get("format") or "xlsx").strip()
    if fmt == "pdf":
        return export_pdf(headers, rows, "اسناد واریز", "deposit_vouchers.pdf")
    return export_xlsx(headers, rows, "اسناد واریز", "deposit_vouchers.xlsx")


@bp.get("/accounting/expenses/export")
@login_required
def expenses_export():
    err = require_finance()
    if err: return err
    cat_lbl = {"land": "خرید زمین", "contractor": "پیمانکار", "admin": "اداری",
               "utility": "خدمات", "other": "سایر"}
    es = db_session.query(ExpenseVoucher).order_by(ExpenseVoucher.number).all()
    rows = [[
        e.number, e.payee, cat_lbl.get(e.category, e.category),
        e.project.name if e.project else "", e.amount, e.spend_date,
        e.reference_no or "", "ابطال" if e.status == "voided" else "فعال",
        e.description or "",
    ] for e in es]
    headers = ["شماره سند", "دریافت‌کننده", "دسته", "پروژه", "مبلغ (ریال)", "تاریخ",
               "شماره پیگیری", "وضعیت", "شرح"]
    fmt = (request.args.get("format") or "xlsx").strip()
    if fmt == "pdf":
        return export_pdf(headers, rows, "هزینه‌ها", "expense_vouchers.pdf")
    return export_xlsx(headers, rows, "هزینه‌ها", "expense_vouchers.xlsx")
